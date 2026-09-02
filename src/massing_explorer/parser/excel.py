from __future__ import annotations

import csv
from pathlib import Path

import openpyxl

from .tabular import parse_tabular_rows


def _read_excel_rows(path: Path) -> tuple[list[list], str]:
    wb = openpyxl.load_workbook(path, data_only=True)
    # Prefer sheets whose name suggests program data
    sheet_names = wb.sheetnames
    preferred = [
        n
        for n in sheet_names
        if any(k in n.lower() for k in ("space", "program", "summary", "schedule", "room"))
    ]
    candidates = preferred or sheet_names

    best_rows: list[list] = []
    best_sheet = candidates[0]
    best_room_count = -1

    for name in candidates:
        ws = wb[name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(list(row))
        try:
            rooms, _, _ = parse_tabular_rows(rows, source_file=str(path))
            if len(rooms) > best_room_count:
                best_room_count = len(rooms)
                best_rows = rows
                best_sheet = name
        except ValueError:
            continue

    if best_room_count < 0:
        ws = wb[best_sheet]
        best_rows = [list(row) for row in ws.iter_rows(values_only=True)]

    return best_rows, best_sheet


def _read_csv_rows(path: Path) -> list[list]:
    rows: list[list] = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        for row in reader:
            rows.append(row)
    return rows


def load_rows_from_file(path: str | Path) -> tuple[list[list], dict]:
    file_path = Path(path)
    meta: dict = {"file": str(file_path), "format": file_path.suffix.lower()}

    if file_path.suffix.lower() in (".xlsx", ".xlsm", ".xltx"):
        rows, sheet = _read_excel_rows(file_path)
        meta["sheet"] = sheet
    elif file_path.suffix.lower() == ".csv":
        rows = _read_csv_rows(file_path)
    else:
        raise ValueError(f"Unsupported file format: {file_path.suffix}")

    return rows, meta

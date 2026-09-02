"""Validate parser against user-provided test Excel files."""

from __future__ import annotations

from pathlib import Path

import openpyxl

from massing_explorer.load import load_program_file

CONFIG = Path(__file__).resolve().parents[1] / "config" / "project.example.yaml"

FILES = {
    "TEST (MSBA variant)": Path(
        r"c:\Users\tu\Downloads\Underwood_Elementary_Space_Summary_TEST.xlsx"
    ),
    "Structured": Path(
        r"c:\Users\tu\Downloads\Underwood_Elementary_Space_Summary_TEST_Structured.xlsx"
    ),
    "Shuffled": Path(
        r"c:\Users\tu\Downloads\Underwood_Elementary_Shuffled_Program_Test.xlsx"
    ),
    "Adversarial": Path(
        r"c:\Users\tu\Downloads\Underwood_Elementary_Adversarial_Category_Pairing_Test.xlsx"
    ),
}

SUMMARY_SHEETS = {
    "Structured": "Summary",
    "Shuffled": "Summary",
    "Adversarial": "Assigned Category Summary",
}


def read_summary_totals(path: Path, sheet_name: str) -> dict[str, float]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name]
    hdr_row = 4
    for r in range(1, 8):
        if ws.cell(r, 1).value and "PROGRAM CATEGORY" in str(ws.cell(r, 1).value):
            hdr_row = r
            break
    expected: dict[str, float] = {}
    for r in range(hdr_row + 1, ws.max_row + 1):
        cat = ws.cell(r, 1).value
        area = ws.cell(r, 2).value
        if cat and isinstance(area, (int, float)):
            expected[str(cat).strip()] = float(area)
    return expected


def main() -> None:
    for label, path in FILES.items():
        print("=" * 72)
        print(label)
        if not path.exists():
            print(f"  MISSING: {path}")
            continue

        study = load_program_file(path, config_path=CONFIG)
        nfa = study.totals["nfa_sf"]
        print(f"  Rooms: {len(study.rooms)}")
        print(f"  Departments: {len(study.departments)}")
        print(f"  NFA total: {nfa:,.0f} SF")
        print(f"  Sheet used: {study.metadata.get('sheet')}")
        print(f"  Column map: {study.metadata.get('column_map')}")

        errors = [v for v in study.verification if not v.passed and v.severity == "error"]
        warnings = [v for v in study.verification if not v.passed and v.severity == "warning"]
        print(f"  Verification: {len(errors)} errors, {len(warnings)} warnings")

        for v in study.verification:
            status = "OK" if v.passed else v.severity.upper()
            print(f"    [{status}] {v.message}")

        summary_key = label.split()[0]
        if summary_key in SUMMARY_SHEETS:
            expected = read_summary_totals(path, SUMMARY_SHEETS[summary_key])
            print("  --- vs summary sheet ---")
            all_ok = True
            for dept in sorted(study.departments, key=lambda d: -d.nfa_sf):
                exp = expected.get(dept.name)
                if exp is None:
                    print(f"    [??] {dept.name}: {dept.nfa_sf:,.0f} (missing from summary)")
                    all_ok = False
                else:
                    ok = abs(dept.nfa_sf - exp) < 1
                    if not ok:
                        all_ok = False
                    mark = "OK" if ok else "FAIL"
                    print(f"    [{mark}] {dept.name}: {dept.nfa_sf:,.0f} vs {exp:,.0f}")
            exp_total = sum(expected.values())
            total_ok = abs(nfa - exp_total) < 1
            print(f"    Total: {nfa:,.0f} vs {exp_total:,.0f} -> {'OK' if total_ok else 'FAIL'}")
            print(f"  RESULT: {'PASS' if all_ok and total_ok else 'FAIL'}")


if __name__ == "__main__":
    main()
